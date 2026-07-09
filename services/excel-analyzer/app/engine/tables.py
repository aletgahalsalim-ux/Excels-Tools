"""Detected Table extraction.

Algorithm (doc phase 6):
1. Materialize the sheet's value grid; merged cells propagate their top-left
   value to every covered cell so region detection is not broken by merges.
2. Find connected regions of non-empty cells (4-neighbour flood fill).
3. Keep regions of at least 2x2 as Detected Tables.
4. Header detection heuristic: the first region row is a header when it is
   mostly text AND the rows beneath it are more numeric than the first row.
"""

from __future__ import annotations

import datetime
from dataclasses import dataclass
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MIN_TABLE_ROWS = 2
MIN_TABLE_COLS = 2


@dataclass
class DetectedTable:
    sheet_name: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int
    header_row: int | None
    headers: list[str]
    rows: list[list[Any]]
    numeric_ratio: float
    row_count: int
    truncated: bool

    @property
    def range_a1(self) -> str:
        return (
            f"{get_column_letter(self.min_col)}{self.min_row}:"
            f"{get_column_letter(self.max_col)}{self.max_row}"
        )

    @property
    def table_id(self) -> str:
        return f"{self.sheet_name}!{self.range_a1}"

    def to_json(self) -> dict[str, Any]:
        return {
            "id": self.table_id,
            "sheetName": self.sheet_name,
            "range": self.range_a1,
            "headerRow": self.header_row,
            "headers": self.headers,
            "rowCount": self.row_count,
            "columnCount": self.max_col - self.min_col + 1,
            "numericRatio": round(self.numeric_ratio, 4),
            "rows": self.rows,
            "truncated": self.truncated,
        }


def _json_value(value: Any) -> Any:
    """Convert a cell value into a JSON-safe value."""
    if value is None or isinstance(value, (int, float, str, bool)):
        return value
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def build_value_grid(ws: Worksheet) -> dict[tuple[int, int], Any]:
    """Grid of non-empty cell values, with merged ranges filled from their anchor."""
    grid: dict[tuple[int, int], Any] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.value != "":
                grid[(cell.row, cell.column)] = cell.value

    for merged in ws.merged_cells.ranges:
        anchor = grid.get((merged.min_row, merged.min_col))
        if anchor is None:
            continue
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                grid.setdefault((r, c), anchor)
    return grid


def _find_regions(grid: dict[tuple[int, int], Any]) -> list[tuple[int, int, int, int]]:
    """Connected regions (min_row, min_col, max_row, max_col) via flood fill."""
    unvisited = set(grid.keys())
    regions: list[tuple[int, int, int, int]] = []
    while unvisited:
        start = unvisited.pop()
        stack = [start]
        min_r = max_r = start[0]
        min_c = max_c = start[1]
        while stack:
            r, c = stack.pop()
            min_r, max_r = min(min_r, r), max(max_r, r)
            min_c, max_c = min(min_c, c), max(max_c, c)
            for nb in ((r - 1, c), (r + 1, c), (r, c - 1), (r, c + 1)):
                if nb in unvisited:
                    unvisited.discard(nb)
                    stack.append(nb)
        regions.append((min_r, min_c, max_r, max_c))
    return regions


def _is_numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _detect_header(
    grid: dict[tuple[int, int], Any], min_r: int, min_c: int, max_r: int, max_c: int
) -> bool:
    """First row is a header when mostly text and the body is more numeric."""
    first = [grid.get((min_r, c)) for c in range(min_c, max_c + 1)]
    present = [v for v in first if v is not None]
    if not present:
        return False
    text_ratio = sum(isinstance(v, str) for v in present) / len(present)

    body_cells = [
        grid[(r, c)]
        for r in range(min_r + 1, max_r + 1)
        for c in range(min_c, max_c + 1)
        if (r, c) in grid
    ]
    if not body_cells:
        return False
    body_numeric = sum(_is_numeric(v) for v in body_cells) / len(body_cells)
    first_numeric = sum(_is_numeric(v) for v in present) / len(present)
    return text_ratio >= 0.6 and body_numeric > first_numeric


def detect_tables(ws: Worksheet, max_rows: int = 200) -> list[DetectedTable]:
    grid = build_value_grid(ws)
    tables: list[DetectedTable] = []

    for min_r, min_c, max_r, max_c in _find_regions(grid):
        if max_r - min_r + 1 < MIN_TABLE_ROWS or max_c - min_c + 1 < MIN_TABLE_COLS:
            continue

        has_header = _detect_header(grid, min_r, min_c, max_r, max_c)
        if has_header:
            headers = [
                str(grid.get((min_r, c), f"Column {i + 1}"))
                for i, c in enumerate(range(min_c, max_c + 1))
            ]
            body_start = min_r + 1
            header_row: int | None = min_r
        else:
            headers = [f"Column {i + 1}" for i in range(max_c - min_c + 1)]
            body_start = min_r
            header_row = None

        all_rows = [
            [_json_value(grid.get((r, c))) for c in range(min_c, max_c + 1)]
            for r in range(body_start, max_r + 1)
        ]
        truncated = len(all_rows) > max_rows

        data_cells = [v for row in all_rows for v in row if v is not None]
        numeric_ratio = (
            sum(isinstance(v, (int, float)) and not isinstance(v, bool) for v in data_cells)
            / len(data_cells)
            if data_cells
            else 0.0
        )

        tables.append(
            DetectedTable(
                sheet_name=ws.title,
                min_row=min_r,
                min_col=min_c,
                max_row=max_r,
                max_col=max_c,
                header_row=header_row,
                headers=headers,
                rows=all_rows[:max_rows],
                numeric_ratio=numeric_ratio,
                row_count=len(all_rows),
                truncated=truncated,
            )
        )

    tables.sort(key=lambda t: (t.min_row, t.min_col))
    return tables
