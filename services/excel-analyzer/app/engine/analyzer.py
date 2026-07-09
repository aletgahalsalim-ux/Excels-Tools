"""Workbook analyzer — orchestrates table detection, formula analysis and
anomaly detection into a single AnalysisResult JSON (contract:
packages/shared/schemas/analysis-result.schema.json).
"""

from __future__ import annotations

from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .formulas import detect_anomalies, extract_formulas
from .tables import build_value_grid, detect_tables

ENGINE_VERSION = "0.1.0"


def analyze_workbook(path: str, file_name: str, max_rows_per_table: int = 200) -> dict[str, Any]:
    # Two loads: cached values for table data, raw formulas for dependency analysis.
    wb_values = load_workbook(path, data_only=True)
    wb_formulas = load_workbook(path, data_only=False)

    sheets_json: list[dict[str, Any]] = []
    total_cells = 0
    total_tables = 0
    value_grids: dict[str, dict[str, Any]] = {}

    for idx, ws in enumerate(wb_values.worksheets):
        grid = build_value_grid(ws)
        total_cells += len(grid)
        value_grids[ws.title] = {f"{get_column_letter(c)}{r}": v for (r, c), v in grid.items()}

        tables = detect_tables(ws, max_rows=max_rows_per_table)
        total_tables += len(tables)

        sheets_json.append(
            {
                "name": ws.title,
                "index": idx,
                "visible": ws.sheet_state == "visible",
                "maxRow": ws.max_row or 0,
                "maxColumn": ws.max_column or 0,
                "mergedCells": [str(r) for r in ws.merged_cells.ranges],
                "detectedTables": [t.to_json() for t in tables],
            }
        )

    formulas = extract_formulas(wb_formulas)
    anomalies = detect_anomalies(formulas, value_grids)

    named_ranges = [
        {"name": name, "target": str(dn.value)}
        for name, dn in wb_values.defined_names.items()
    ]

    return {
        "engineVersion": ENGINE_VERSION,
        "workbook": {
            "fileName": file_name,
            "sheetCount": len(wb_values.worksheets),
            "namedRanges": named_ranges,
        },
        "sheets": sheets_json,
        "formulas": [f.to_json() for f in formulas],
        "anomalies": anomalies,
        "stats": {
            "totalCells": total_cells,
            "totalFormulas": len(formulas),
            "totalDetectedTables": total_tables,
        },
    }
