from openpyxl import load_workbook

from app.engine.tables import detect_tables


def test_detects_two_separate_tables_on_one_sheet(balanced_model):
    wb = load_workbook(balanced_model, data_only=True)
    ws = wb["Income Statement"]
    tables = detect_tables(ws)
    assert len(tables) == 2
    ranges = {t.range_a1 for t in tables}
    assert "A1:B8" in ranges or "A1:B6" in ranges  # main table (B8 formula joins via col B)
    assert any(t.min_col == 4 for t in tables)  # quarter table starts at column D


def test_header_detection(balanced_model):
    wb = load_workbook(balanced_model, data_only=True)
    tables = detect_tables(wb["Balance Sheet"])
    assert len(tables) == 1
    t = tables[0]
    assert t.header_row == 1
    assert t.headers == ["Item", "2024", "2025"]
    assert t.row_count == 3
    assert t.rows[0] == ["Total Assets", 1000000, 1250000]


def test_numeric_ratio(balanced_model):
    wb = load_workbook(balanced_model, data_only=True)
    t = detect_tables(wb["Balance Sheet"])[0]
    # body: 3 label cells + 6 numeric cells
    assert 0.6 < t.numeric_ratio <= 0.7


def test_merged_cells_do_not_split_table(broken_model):
    wb = load_workbook(broken_model, data_only=True)
    tables = detect_tables(wb["Model"])
    main = max(tables, key=lambda t: t.row_count * (t.max_col - t.min_col + 1))
    # merged A1:B1 title is connected to the table under it
    assert main.min_row == 1


def test_truncation_cap(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Value"])
    for i in range(50):
        ws.append([f"row{i}", i])
    p = tmp_path / "long.xlsx"
    wb.save(p)

    ws2 = load_workbook(p, data_only=True).active
    t = detect_tables(ws2, max_rows=10)[0]
    assert t.truncated is True
    assert len(t.rows) == 10
    assert t.row_count == 50
