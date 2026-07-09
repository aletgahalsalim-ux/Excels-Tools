"""Reference fixture workbooks, generated programmatically (doc phase 11):
- balanced_model.xlsx : clean budget/balance-sheet workbook with headers
- broken_model.xlsx   : deliberate errors — circular ref, hardcoded number,
                        #REF!, dead formula, merged header cells
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook


@pytest.fixture()
def balanced_model(tmp_path: Path) -> Path:
    wb = Workbook()

    ws = wb.active
    ws.title = "Balance Sheet"
    rows = [
        ["Item", "2024", "2025"],
        ["Total Assets", 1000000, 1250000],
        ["Total Liabilities", 400000, 500000],
        ["Equity", 600000, 750000],
    ]
    for r in rows:
        ws.append(r)

    ws2 = wb.create_sheet("Income Statement")
    rows2 = [
        ["Item", "Amount"],
        ["Revenue", 800000],
        ["COGS", 300000],
        ["Gross Profit", 500000],
        ["OPEX", 200000],
        ["Net Income", 300000],
    ]
    for r in rows2:
        ws2.append(r)

    # a second, separate table on the same sheet (blank row/col gap)
    ws2["D1"] = "Quarter"
    ws2["E1"] = "Sales"
    for i, (q, v) in enumerate([("Q1", 150000), ("Q2", 200000), ("Q3", 220000), ("Q4", 230000)]):
        ws2.cell(row=i + 2, column=4, value=q)
        ws2.cell(row=i + 2, column=5, value=v)

    # cross-sheet formula
    ws2["B8"] = "='Balance Sheet'!B2"

    wb.defined_names.add(_defined_name("TotalAssets", "'Balance Sheet'!$B$2"))

    path = tmp_path / "balanced_model.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def broken_model(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"

    # merged header over two columns
    ws["A1"] = "Financial Model"
    ws.merge_cells("A1:B1")

    ws["A2"] = "Metric"
    ws["B2"] = "Value"
    ws["A3"] = "Revenue"
    ws["B3"] = 500000
    ws["A4"] = "Costs"
    ws["B4"] = 600000
    ws["A5"] = "Margin"
    # hardcoded number inside formula
    ws["B5"] = "=B3-B4+12345"
    # circular reference pair
    ws["C3"] = "=C4"
    ws["C4"] = "=C3"
    # broken reference
    ws["C5"] = "=#REF!+B3"
    # dead formula: references empty cells only
    ws["C6"] = "=Z100+Z101"

    path = tmp_path / "broken_model.xlsx"
    wb.save(path)
    return path


def _defined_name(name: str, target: str):
    from openpyxl.workbook.defined_name import DefinedName

    return DefinedName(name, attr_text=target)
