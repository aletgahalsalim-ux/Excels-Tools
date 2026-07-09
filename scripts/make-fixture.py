#!/usr/bin/env python3
"""Generate the e2e fixture workbook: a financial model with deliberate errors
(unbalanced balance sheet, circular reference, hardcoded number in formula)."""

import sys

from openpyxl import Workbook

out = sys.argv[1] if len(sys.argv) > 1 else "fixture-model.xlsx"

wb = Workbook()

bs = wb.active
bs.title = "Balance Sheet"
for row in [
    ["البند", "Item", "2025"],
    ["إجمالي الأصول", "Total Assets", 1_000_000],
    ["إجمالي الخصوم", "Total Liabilities", 400_000],
    ["حقوق الملكية", "Equity", 500_000],  # planted error: 400k + 500k != 1000k
]:
    bs.append(row)

inc = wb.create_sheet("Income Statement")
for row in [
    ["Item", "Amount"],
    ["Revenue", 800_000],
    ["Total Costs", 500_000],
    ["Net Income", 300_000],
]:
    inc.append(row)

model = wb.create_sheet("Model")
model["A1"] = "Calc"
model["B1"] = "Value"
model["A2"] = "Forecast"
model["B2"] = "=B3*1.15"  # hardcoded growth factor
model["A3"] = "Base"
model["B3"] = "=B2"  # circular with B2

wb.save(out)
print(f"fixture written to {out}")
